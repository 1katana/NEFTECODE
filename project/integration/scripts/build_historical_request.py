"""Build an observation-only API request from time-aligned source records.

The LIMS result is written only to the evidence file: it was not available to
the decision system at the sample timestamp. No blending recipe is invented.
"""

from __future__ import annotations

import argparse
import json
from datetime import UTC
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from quality_agent.feature_pipeline import FeatureSchema
from quality_agent.runtime import RuntimeConfig
from reliability_agent.config import RAW_FEATURES
from reliability_agent.data import read_telemetry, snapshot

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_AT = "2026-05-15 10:00:00"
DEFAULT_OUTPUT = ROOT / "integration" / "examples" / "historical_20260515_1000.json"


def _read_tag_window(start: pd.Timestamp, end: pd.Timestamp) -> pd.DataFrame:
    runtime = RuntimeConfig.from_yaml(ROOT / "Quality_neftecode/configs/runtime_shadow.yaml")
    schema = FeatureSchema.load(runtime.primary_artifacts / "feature_schema.json")
    frames = []
    for source, name in (("avt", "avt_tags.csv"), ("hyd", "242000_tags.csv")):
        tags = sorted(schema.required_signals[source])
        matches = []
        for chunk in pd.read_csv(
            ROOT / "data/data" / name,
            usecols=["date", *tags],
            chunksize=20_000,
        ):
            selected = chunk.loc[chunk["date"].between(str(start), str(end))]
            if not selected.empty:
                matches.append(selected)
        if not matches:
            raise ValueError(f"No {source} telemetry between {start} and {end}")
        frame = pd.concat(matches).set_index("date")
        if frame.index.has_duplicates:
            raise ValueError(f"Duplicate timestamps in {name}")
        frames.append(frame.rename(columns={tag: f"{source}_{tag.lower()}" for tag in tags}))
    combined = pd.concat(frames, axis=1)
    expected = pd.date_range(start, end, freq="10min")
    if combined.index.tolist() != [str(item) for item in expected]:
        raise ValueError("AVT and hydrotreater data must cover the exact 10-minute window")
    if not set(RAW_FEATURES).issubset(combined.columns):
        raise ValueError("Quality feature schema does not include all Reliability inputs")
    return combined


def _read_pak_window(start: pd.Timestamp, end: pd.Timestamp) -> dict[str, dict[str, float]]:
    workbook = load_workbook(
        ROOT / "Выгрузка ПАК 01.01.2023 - н.в_.xlsx",
        read_only=True,
        data_only=True,
    )
    readings: dict[str, dict[str, float]] = {}
    try:
        for sulfur_time, sulfur, _, density_time, density in workbook.active.iter_rows(
            min_row=3, min_col=1, max_col=5, values_only=True
        ):
            for timestamp, key, value in (
                (sulfur_time, "mg_sulfur", sulfur),
                (density_time, "pak_d15", density),
            ):
                if timestamp is None or not start <= pd.Timestamp(timestamp) <= end:
                    continue
                if value is None or not pd.notna(value):
                    continue
                readings.setdefault(str(timestamp), {})[key] = float(value)
    finally:
        workbook.close()
    return readings


def _reference_lims(at: pd.Timestamp) -> dict[str, object] | None:
    path = ROOT / "Quality_neftecode/data/processed/quality_training_lims_samples.csv"
    samples = pd.read_csv(
        path,
        usecols=["sample_id", "sample_timestamp", "target_mg_sulfur", "calendar_split"],
    )
    matching = samples.loc[samples["sample_timestamp"] == str(at)]
    if matching.empty:
        return None
    if len(matching) != 1:
        raise ValueError(f"Expected one LIMS sample at {at}, got {len(matching)}")
    row = matching.iloc[0]
    return {
        "sample_id": str(row["sample_id"]),
        "sample_timestamp": str(row["sample_timestamp"]),
        "sulfur_mg_kg": float(row["target_mg_sulfur"]),
        "calendar_split": str(row["calendar_split"]),
        "used_as_model_input": False,
    }


def build(at: pd.Timestamp) -> tuple[dict[str, object], dict[str, object]]:
    if at.tzinfo is not None:
        raise ValueError("Source files use local naive timestamps; --at must have no timezone")
    if at != at.floor("10min"):
        raise ValueError("--at must be on a 10-minute telemetry boundary")
    start = at - pd.Timedelta(hours=6)
    telemetry = _read_tag_window(start, at)
    pak = _read_pak_window(start, at)
    expected = [str(item) for item in pd.date_range(start, at, freq="10min")]
    missing_pak = [time for time in expected if set(pak.get(time, {})) != {"mg_sulfur", "pak_d15"}]
    if missing_pak:
        raise ValueError(
            f"PAK sulfur and D15 must be present at every timestamp: {missing_pak[:3]}"
        )
    reliability_values, flatline_ages = read_telemetry()
    states = []
    for time in expected:
        row = telemetry.loc[time]
        original = snapshot(reliability_values, flatline_ages, at=time)
        measured = {tag: float(value) if pd.notna(value) else None for tag, value in row.items()}
        missing = sorted(tag for tag, value in measured.items() if value is None)
        quality_flags = dict(original["data_quality"])
        quality_flags.update({"missing": missing, "stale": [], "outlier": []})
        controls = {tag: measured[tag] for tag in ("hyd_t6", "hyd_f9", "hyd_p13")}
        if any(value is None for value in controls.values()):
            raise ValueError(f"Missing observed control signal at {time}")
        states.append(
            {
                "timestamp": pd.Timestamp(time).tz_localize(UTC).isoformat(),
                "quality": {**pak[time], "source": "pak"},
                "telemetry": measured,
                "data_quality": quality_flags,
                "data_confidence": "HIGH" if not missing else "MEDIUM",
                "current_controls": controls,
            }
        )
    request: dict[str, object] = {
        "run_id": f"historical-{at:%Y%m%d-%H%M}",
        "input_source": "HISTORICAL",
        "process_state": states[-1],
        "history": states[:-1],
    }
    evidence: dict[str, object] = {
        "anchor_timestamp_local": str(at),
        "observation_count": len(states),
        "input_source": "HISTORICAL",
        "sources": {
            "telemetry": ["data/data/avt_tags.csv", "data/data/242000_tags.csv"],
            "pak": "Выгрузка ПАК 01.01.2023 - н.в_.xlsx",
            "lims_reference": "Quality_neftecode/data/processed/quality_training_lims_samples.csv",
        },
        "reference_lims": _reference_lims(at),
        "limitations": [
            "No observed commercial blending recipe or confirmed blend component properties",
            "LIMS reference is withheld from the API request to avoid temporal leakage",
            "Source timestamps have no timezone; UTC is assigned only for API parsing",
        ],
    }
    return request, evidence


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--at", default=DEFAULT_AT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    request, evidence = build(pd.Timestamp(args.at))
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(request, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8"
    )
    evidence_path = output.with_name(output.stem + "_evidence.json")
    evidence_path.write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(output)
    print(evidence_path)


if __name__ == "__main__":
    main()
